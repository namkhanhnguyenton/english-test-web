import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from models import db, User, Level, Test, Question, Submission, SubmissionAnswer
from ppt_converter import convert_ppt_to_slides

def create_sample_pptx(filepath, title, questions_data):
    """
    Creates a styled PPTX file with slide-by-slide questions.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # Blank slide

    # Color palette
    NAVY = RGBColor(26, 43, 76)
    WHITE = RGBColor(255, 255, 255)
    LIGHT_BG = RGBColor(245, 247, 250)
    DARK_TEXT = RGBColor(30, 41, 59)
    GRAY_TEXT = RGBColor(100, 116, 139)
    ACCENT_BLUE = RGBColor(37, 99, 235)
    CARD_BG = RGBColor(255, 255, 255)

    # 1. Slide Introduction
    intro_slide = prs.slides.add_slide(blank_layout)
    bg = intro_slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height) # Rectangle
    bg.fill.solid()
    bg.fill.fore_color.rgb = NAVY
    bg.line.color.rgb = NAVY

    tx_box = intro_slide.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.33), Inches(3.0))
    tf = tx_box.text_frame
    tf.word_wrap = True
    
    p = tf.paragraphs[0]
    p.text = "ENGLISH PROFICIENCY TEST"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(147, 197, 253)
    p.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = title
    p2.font.size = Pt(40)
    p2.font.bold = True
    p2.font.color.rgb = WHITE
    p2.alignment = PP_ALIGN.CENTER

    p3 = tf.add_paragraph()
    p3.text = f"Số lượng: {len(questions_data)} câu hỏi | Thời gian: Theo quy định\nHãy quan sát câu hỏi và tích chọn đáp án tương ứng trên phiếu làm bài."
    p3.font.size = Pt(20)
    p3.font.color.rgb = RGBColor(226, 232, 240)
    p3.alignment = PP_ALIGN.CENTER

    # 2. Slide for each question
    for q_idx, q in enumerate(questions_data, start=1):
        slide = prs.slides.add_slide(blank_layout)

        # Background
        bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = LIGHT_BG
        bg.line.color.rgb = LIGHT_BG

        # Top Header Bar
        header = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(1.2))
        header.fill.solid()
        header.fill.fore_color.rgb = NAVY
        header.line.color.rgb = NAVY

        # Header Title
        h_box = slide.shapes.add_textbox(Inches(1.0), Inches(0.25), Inches(8.0), Inches(0.7))
        h_tf = h_box.text_frame
        hp = h_tf.paragraphs[0]
        hp.text = f"{title.upper()} — CÂU HỎI {q_idx}/{len(questions_data)}"
        hp.font.size = Pt(22)
        hp.font.bold = True
        hp.font.color.rgb = WHITE

        # Question Main Box (White card)
        card = slide.shapes.add_shape(1, Inches(1.0), Inches(1.6), Inches(11.33), Inches(5.2))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = RGBColor(203, 213, 225)
        card.line.width = Pt(1.5)

        # Question Text
        q_box = slide.shapes.add_textbox(Inches(1.5), Inches(1.9), Inches(10.33), Inches(1.5))
        q_tf = q_box.text_frame
        q_tf.word_wrap = True
        qp = q_tf.paragraphs[0]
        qp.text = f"Question {q_idx}: {q['text']}"
        qp.font.size = Pt(30)
        qp.font.bold = True
        qp.font.color.rgb = DARK_TEXT

        # Options Container
        opt_box = slide.shapes.add_textbox(Inches(1.5), Inches(3.6), Inches(10.33), Inches(2.8))
        opt_tf = opt_box.text_frame
        opt_tf.word_wrap = True

        for opt_key, opt_text in q['options'].items():
            p_opt = opt_tf.add_paragraph() if opt_tf.paragraphs[0].text else opt_tf.paragraphs[0]
            p_opt.text = f"{opt_key}. {opt_text}"
            p_opt.font.size = Pt(26)
            p_opt.font.bold = False
            p_opt.font.color.rgb = DARK_TEXT
            p_opt.space_after = Pt(14)

    os.makedirs(os.path.dirname(filepath), exist_ok=True)
    prs.save(filepath)
    return filepath


def seed_database(app):
    with app.app_context():
        db.create_all()

        # 1. Admin User
        admin = User.query.filter_by(username='admin').first()
        if not admin:
            admin = User(
                username='admin',
                full_name='Quản trị viên Hệ thống',
                email='admin@testenglish.vn',
                role='admin'
            )
            admin.set_password('admin123')
            db.session.add(admin)

        # 1.2 Teacher User (Mẫu)
        teacher = User.query.filter_by(username='giaovien').first()
        if not teacher:
            teacher = User(
                username='giaovien',
                full_name='Cô Nguyễn Thu Hà (Giáo viên)',
                email='giaovien@testenglish.vn',
                role='teacher'
            )
            teacher.set_password('giaovien123')
            db.session.add(teacher)

        # Demo regular user
        user = User.query.filter_by(username='hocvien1').first()
        if not user:
            user = User(
                username='hocvien1',
                full_name='Nguyễn Văn An',
                email='an.nguyen@example.com',
                role='user'
            )
            user.set_password('123456')
            db.session.add(user)

        db.session.commit()

        # 2. Levels
        levels_data = [
            {
                'code': 'A1',
                'name': 'Trình độ A1 - Căn bản (Beginner)',
                'description': 'Dành cho người mới bắt đầu hoặc mất gốc. Kiểm tra ngữ pháp, thì cơ bản và từ vựng thông dụng.',
                'badge_color': 'emerald',
                'icon_name': 'book-open',
                'order_num': 1
            },
            {
                'code': 'A2',
                'name': 'Trình độ A2 - Sơ cấp (Elementary)',
                'description': 'Kiểm tra khả năng giao tiếp đơn giản, so sánh, câu điều kiện loại 1 và vốn từ vựng mở rộng.',
                'badge_color': 'blue',
                'icon_name': 'sparkles',
                'order_num': 2
            },
            {
                'code': 'B1',
                'name': 'Trình độ B1 - Trung cấp (Intermediate)',
                'description': 'Kiểm tra ngữ pháp chuyên sâu, mệnh đề quan hệ, câu bị động, ngữ đồng nghĩa và đọc hiểu.',
                'badge_color': 'indigo',
                'icon_name': 'award',
                'order_num': 3
            },
            {
                'code': 'B2',
                'name': 'Trình độ B2 - Trung cao cấp (Upper Intermediate)',
                'description': 'Đánh giá khả năng hiểu văn bản phức tạp, thành ngữ (idioms), collocations và phân tích ngữ cảnh.',
                'badge_color': 'purple',
                'icon_name': 'target',
                'order_num': 4
            },
            {
                'code': 'IELTS',
                'name': 'Luyện thi IELTS & TOEIC',
                'description': 'Đề thi chuẩn hóa định dạng quốc tế, từ vựng học thuật (academic), cấu trúc câu nâng cao.',
                'badge_color': 'amber',
                'icon_name': 'zap',
                'order_num': 5
            }
        ]

        level_objs = {}
        for l_data in levels_data:
            lvl = Level.query.filter_by(code=l_data['code']).first()
            if not lvl:
                lvl = Level(**l_data)
                db.session.add(lvl)
                db.session.commit()
            level_objs[l_data['code']] = lvl

        # 3. Create Sample Tests if not exist
        if Test.query.count() == 0:
            print("Creating sample PPT tests and slides...")

            # --- TEST 1: Level A1 - Đề số 1 ---
            q_data_a1_1 = [
                {'text': 'She ___ to work by bus every morning.', 'options': {'A': 'go', 'B': 'goes', 'C': 'going', 'D': 'gone'}, 'ans': 'B', 'exp': 'Chủ ngữ ngôi thứ ba số ít "She" đi với động từ thêm -es "goes".'},
                {'text': 'There ___ twenty students in the classroom.', 'options': {'A': 'is', 'B': 'are', 'C': 'be', 'D': 'was'}, 'ans': 'B', 'exp': 'Danh từ số nhiều "twenty students" dùng "There are".'},
                {'text': 'What is the opposite of "Cheap"?', 'options': {'A': 'Free', 'B': 'Expensive', 'C': 'Easy', 'D': 'Slow'}, 'ans': 'B', 'exp': '"Cheap" (rẻ) trái nghĩa với "Expensive" (đắt).'},
                {'text': 'I do not have ___ money in my wallet.', 'options': {'A': 'any', 'B': 'some', 'C': 'many', 'D': 'a few'}, 'ans': 'A', 'exp': 'Trong câu phủ định với danh từ không đếm được dùng "any".'},
                {'text': 'Where ___ you yesterday afternoon?', 'options': {'A': 'was', 'B': 'were', 'C': 'are', 'D': 'did'}, 'ans': 'B', 'exp': 'Chủ ngữ "you" ở quá khứ đơn đi với to be "were".'},
                {'text': 'He is very good ___ playing chess.', 'options': {'A': 'in', 'B': 'at', 'C': 'on', 'D': 'with'}, 'ans': 'B', 'exp': 'Cấu trúc "good at something" nghĩa là giỏi việc gì.'},
                {'text': 'Listen! The phone ___. Can you answer it?', 'options': {'A': 'rings', 'B': 'rang', 'C': 'is ringing', 'D': 'has rung'}, 'ans': 'C', 'exp': 'Dấu hiệu "Listen!" chỉ hành động đang xảy ra tại thời điểm nói (Hiện tại tiếp diễn).'},
                {'text': 'Last summer, my family ___ to Da Nang for vacation.', 'options': {'A': 'go', 'B': 'went', 'C': 'gone', 'D': 'going'}, 'ans': 'B', 'exp': '"Last summer" là dấu hiệu của thì Quá khứ đơn -> "went".'},
                {'text': 'This car is much ___ than that old one.', 'options': {'A': 'fast', 'B': 'faster', 'C': 'fastest', 'D': 'more fast'}, 'ans': 'B', 'exp': 'So sánh hơn của tính từ ngắn "fast" là "faster than".'},
                {'text': 'Would you like ___ apple or an orange?', 'options': {'A': 'a', 'B': 'an', 'C': 'the', 'D': 'some'}, 'ans': 'B', 'exp': 'Từ "apple" bắt đầu bằng nguyên âm [æ] nên dùng mạo từ "an".'}
            ]

            upload_test1_dir = os.path.join(app.root_path, 'static', 'uploads', 'tests', 'test_1')
            os.makedirs(upload_test1_dir, exist_ok=True)
            pptx_test1_path = os.path.join(upload_test1_dir, 'A1_Grammar_Test_01.pptx')
            create_sample_pptx(pptx_test1_path, "A1 Grammar & Vocabulary - Đề 01", q_data_a1_1)

            # Convert to slides
            slides_folder_1 = os.path.join(upload_test1_dir, 'slides')
            slide_files_1, total_slides_1 = convert_ppt_to_slides(pptx_test1_path, slides_folder_1)

            test_1 = Test(
                level_id=level_objs['A1'].id,
                title="Đề số 01: Ngữ pháp & Từ vựng Căn bản A1",
                description="Đề thi 10 câu trắc nghiệm kiểm tra thì hiện tại đơn, quá khứ đơn, mạo từ và giới từ căn bản.",
                ppt_filename="A1_Grammar_Test_01.pptx",
                ppt_filepath='uploads/tests/test_1/A1_Grammar_Test_01.pptx',
                slides_folder='uploads/tests/test_1/slides',
                total_slides=total_slides_1,
                time_limit_minutes=15,
                total_questions=len(q_data_a1_1),
                pass_percentage=50,
                is_active=True
            )
            db.session.add(test_1)
            db.session.commit()

            # Add questions for test 1
            for idx, q in enumerate(q_data_a1_1, start=1):
                # Note: slide 1 is intro, slide idx+1 is question slide
                q_obj = Question(
                    test_id=test_1.id,
                    question_number=idx,
                    question_text=q['text'],
                    question_type='single_choice',
                    options_json='["A", "B", "C", "D"]',
                    correct_answer=q['ans'],
                    slide_number=idx + 1 if total_slides_1 >= idx + 1 else 1,
                    points=1.0,
                    explanation=q['exp']
                )
                db.session.add(q_obj)

            # --- TEST 2: Level A1 - Đề số 2 (to demonstrate RANDOM TEST in A1) ---
            q_data_a1_2 = [
                {'text': 'My sister ___ like drinking coffee.', 'options': {'A': 'don\'t', 'B': 'doesn\'t', 'C': 'isn\'t', 'D': 'aren\'t'}, 'ans': 'B', 'exp': 'Chủ ngữ "My sister" số ít phủ định dùng "doesn\'t".'},
                {'text': 'How ___ brothers and sisters do you have?', 'options': {'A': 'much', 'B': 'many', 'C': 'long', 'D': 'often'}, 'ans': 'B', 'exp': 'Danh từ số nhiều đếm được dùng "How many".'},
                {'text': 'I usually wake up ___ 6:30 AM every morning.', 'options': {'A': 'at', 'B': 'on', 'C': 'in', 'D': 'to'}, 'ans': 'A', 'exp': 'Chỉ giờ giấc cụ thể đi với giới từ "at".'},
                {'text': 'The weather today is ___ than yesterday.', 'options': {'A': 'hot', 'B': 'hotter', 'C': 'hottest', 'D': 'more hot'}, 'ans': 'B', 'exp': 'So sánh hơn của "hot" gấp đôi phụ âm: "hotter".'},
                {'text': 'Can you help me ___ this heavy box?', 'options': {'A': 'carrying', 'B': 'carry', 'C': 'carried', 'D': 'to carried'}, 'ans': 'B', 'exp': 'Cấu trúc "help someone (to) do something".'},
                {'text': 'They ___ television when the power went out.', 'options': {'A': 'watched', 'B': 'were watching', 'C': 'are watching', 'D': 'have watched'}, 'ans': 'B', 'exp': 'Hành động đang diễn ra trong quá khứ thì bị hành động khác xen vào.'}
            ]

            upload_test2_dir = os.path.join(app.root_path, 'static', 'uploads', 'tests', 'test_2')
            os.makedirs(upload_test2_dir, exist_ok=True)
            pptx_test2_path = os.path.join(upload_test2_dir, 'A1_Quick_Quiz_02.pptx')
            create_sample_pptx(pptx_test2_path, "A1 Quick Quiz - Đề 02", q_data_a1_2)

            slides_folder_2 = os.path.join(upload_test2_dir, 'slides')
            slide_files_2, total_slides_2 = convert_ppt_to_slides(pptx_test2_path, slides_folder_2)

            test_2 = Test(
                level_id=level_objs['A1'].id,
                title="Đề số 02: Ôn tập Nhanh Ngữ pháp A1",
                description="Đề thi 6 câu trắc nghiệm tổng hợp cấu trúc câu thông dụng.",
                ppt_filename="A1_Quick_Quiz_02.pptx",
                ppt_filepath='uploads/tests/test_2/A1_Quick_Quiz_02.pptx',
                slides_folder='uploads/tests/test_2/slides',
                total_slides=total_slides_2,
                time_limit_minutes=10,
                total_questions=len(q_data_a1_2),
                pass_percentage=50,
                is_active=True
            )
            db.session.add(test_2)
            db.session.commit()

            for idx, q in enumerate(q_data_a1_2, start=1):
                q_obj = Question(
                    test_id=test_2.id,
                    question_number=idx,
                    question_text=q['text'],
                    question_type='single_choice',
                    options_json='["A", "B", "C", "D"]',
                    correct_answer=q['ans'],
                    slide_number=idx + 1 if total_slides_2 >= idx + 1 else 1,
                    points=1.0,
                    explanation=q['exp']
                )
                db.session.add(q_obj)

            # --- TEST 3: Level B1 - Đề số 1 ---
            q_data_b1_1 = [
                {'text': 'If I had enough money, I ___ a new laptop right now.', 'options': {'A': 'will buy', 'B': 'would buy', 'C': 'bought', 'D': 'had bought'}, 'ans': 'B', 'exp': 'Câu điều kiện loại 2: If + S + V2/ed, S + would + V-inf.'},
                {'text': 'The bridge ___ built in 1995 by an international company.', 'options': {'A': 'is', 'B': 'has been', 'C': 'was', 'D': 'were'}, 'ans': 'C', 'exp': 'Bị động quá khứ đơn: was/were + V3/ed.'},
                {'text': 'She avoided ___ about her past experiences.', 'options': {'A': 'to talk', 'B': 'talk', 'C': 'talking', 'D': 'talked'}, 'ans': 'C', 'exp': 'Động từ "avoid" đi kèm V-ing.'},
                {'text': 'The man ___ car was stolen called the police immediately.', 'options': {'A': 'who', 'B': 'whom', 'C': 'whose', 'D': 'which'}, 'ans': 'C', 'exp': 'Đại từ quan hệ chỉ sở hữu dùng "whose".'},
                {'text': 'We have been friends ___ we were in primary school.', 'options': {'A': 'for', 'B': 'since', 'C': 'during', 'D': 'from'}, 'ans': 'B', 'exp': 'Dấu hiệu thì Hiện tại hoàn thành với mốc thời gian dùng "since".'}
            ]

            upload_test3_dir = os.path.join(app.root_path, 'static', 'uploads', 'tests', 'test_3')
            os.makedirs(upload_test3_dir, exist_ok=True)
            pptx_test3_path = os.path.join(upload_test3_dir, 'B1_Intermediate_Test_01.pptx')
            create_sample_pptx(pptx_test3_path, "B1 Intermediate Test - Đề 01", q_data_b1_1)

            slides_folder_3 = os.path.join(upload_test3_dir, 'slides')
            slide_files_3, total_slides_3 = convert_ppt_to_slides(pptx_test3_path, slides_folder_3)

            test_3 = Test(
                level_id=level_objs['B1'].id,
                title="Đề số 01: Đánh giá Năng lực Trung cấp B1",
                description="Đề thi B1 đánh giá câu điều kiện, câu bị động, mệnh đề quan hệ và danh động từ.",
                ppt_filename="B1_Intermediate_Test_01.pptx",
                ppt_filepath='uploads/tests/test_3/B1_Intermediate_Test_01.pptx',
                slides_folder='uploads/tests/test_3/slides',
                total_slides=total_slides_3,
                time_limit_minutes=15,
                total_questions=len(q_data_b1_1),
                pass_percentage=60,
                is_active=True
            )
            db.session.add(test_3)
            db.session.commit()

            for idx, q in enumerate(q_data_b1_1, start=1):
                q_obj = Question(
                    test_id=test_3.id,
                    question_number=idx,
                    question_text=q['text'],
                    question_type='single_choice',
                    options_json='["A", "B", "C", "D"]',
                    correct_answer=q['ans'],
                    slide_number=idx + 1 if total_slides_3 >= idx + 1 else 1,
                    points=1.0,
                    explanation=q['exp']
                )
                db.session.add(q_obj)

            db.session.commit()
            print("Successfully initialized levels and sample tests!")
        else:
            print("Database already has tests. Skipping test seeding.")

